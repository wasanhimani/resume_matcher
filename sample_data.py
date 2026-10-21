"""
Generates realistic sample resumes (PDF + DOCX) and job descriptions (TXT)
so we have data to build and test the parsing/matching pipeline against.

Run once: python generate_sample_data.py
"""
import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document

RESUME_DIR = "data/sample_resumes"
JD_DIR = "data/sample_jds"
os.makedirs(RESUME_DIR, exist_ok=True)
os.makedirs(JD_DIR, exist_ok=True)

styles = getSampleStyleSheet()
heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], spaceBefore=10, spaceAfter=4)
body_style = styles['Normal']


def make_pdf_resume(filename, name, email, phone, summary, skills, experience, education):
    doc = SimpleDocTemplate(os.path.join(RESUME_DIR, filename), pagesize=letter)
    story = [
        Paragraph(name, styles['Title']),
        Paragraph(f"{email} | {phone}", body_style),
        Spacer(1, 10),
        Paragraph("Summary", heading_style),
        Paragraph(summary, body_style),
        Paragraph("Skills", heading_style),
        Paragraph(", ".join(skills), body_style),
        Paragraph("Experience", heading_style),
    ]
    for role, org, dates, desc in experience:
        story.append(Paragraph(f"<b>{role}</b> — {org} ({dates})", body_style))
        story.append(Paragraph(desc, body_style))
        story.append(Spacer(1, 6))
    story.append(Paragraph("Education", heading_style))
    for edu in education:
        story.append(Paragraph(edu, body_style))
    doc.build(story)


def make_docx_resume(filename, name, email, phone, summary, skills, experience, education):
    doc = Document()
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"{email} | {phone}")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(summary)
    doc.add_heading("Skills", level=2)
    doc.add_paragraph(", ".join(skills))
    doc.add_heading("Experience", level=2)
    for role, org, dates, desc in experience:
        doc.add_paragraph(f"{role} — {org} ({dates})", style='Intense Quote')
        doc.add_paragraph(desc)
    doc.add_heading("Education", level=2)
    for edu in education:
        doc.add_paragraph(edu)
    doc.save(os.path.join(RESUME_DIR, filename))


def make_jd(filename, title, company, description, requirements):
    with open(os.path.join(JD_DIR, filename), "w") as f:
        f.write(f"Job Title: {title}\n")
        f.write(f"Company: {company}\n\n")
        f.write("Job Description:\n")
        f.write(description + "\n\n")
        f.write("Requirements:\n")
        for r in requirements:
            f.write(f"- {r}\n")


# ---------------- RESUMES ----------------

make_pdf_resume(
    "priya_sharma_data_scientist.pdf",
    "Priya Sharma", "priya.sharma@email.com", "+91-9876543210",
    "Data Scientist with 3 years of experience building ML models for e-commerce recommendation systems.",
    ["Python", "Pandas", "NumPy", "Scikit-learn", "TensorFlow", "SQL", "Machine Learning", "Data Visualization", "Statistics"],
    [("Data Scientist", "Flipkart", "2022-Present", "Built recommendation engine using collaborative filtering, improved CTR by 18%. Used Python, pandas, scikit-learn."),
     ("ML Intern", "Zomato", "2021-2022", "Worked on churn prediction models using logistic regression and XGBoost.")],
    ["B.Tech in Computer Science, IIT Delhi, 2022"],
)

make_pdf_resume(
    "rohan_verma_backend_developer.pdf",
    "Rohan Verma", "rohan.verma@email.com", "+91-9812345678",
    "Backend Developer with 4 years of experience building scalable microservices in Java and Spring Boot.",
    ["Java", "Spring Boot", "Microservices", "REST APIs", "SQL", "PostgreSQL", "Docker", "Kubernetes", "AWS"],
    [("Backend Developer", "Paytm", "2021-Present", "Designed REST APIs and microservices using Spring Boot, deployed on Kubernetes on AWS."),
     ("Software Engineer", "Infosys", "2020-2021", "Maintained Java based banking applications, wrote SQL queries for reporting.")],
    ["B.Tech in Information Technology, NIT Trichy, 2020"],
)

make_pdf_resume(
    "ananya_singh_frontend_developer.pdf",
    "Ananya Singh", "ananya.singh@email.com", "+91-9900112233",
    "Frontend Developer with 2 years of experience creating responsive web applications using React.",
    ["JavaScript", "TypeScript", "React", "Redux", "HTML", "CSS", "Tailwind CSS", "Git"],
    [("Frontend Developer", "Swiggy", "2023-Present", "Built and maintained customer facing React applications, improved page load time by 30%."),
     ("Web Developer Intern", "Freelance", "2022-2023", "Built landing pages using HTML, CSS and JavaScript for small businesses.")],
    ["B.Tech in Computer Science, VIT Vellore, 2023"],
)

make_docx_resume(
    "karan_mehta_fullstack_developer.docx",
    "Karan Mehta", "karan.mehta@email.com", "+91-9765432109",
    "Full Stack Developer with 3 years of experience building web applications using the MERN stack.",
    ["JavaScript", "Node.js", "Express", "React", "MongoDB", "REST APIs", "AWS", "Git"],
    [("Full Stack Developer", "Razorpay", "2022-Present", "Built and shipped features across the MERN stack, integrated payment APIs, deployed on AWS."),
     ("Software Developer", "TCS", "2021-2022", "Worked on internal tools using Node.js and MongoDB.")],
    ["B.Tech in Computer Engineering, Pune Institute of Computer Technology, 2021"],
)

make_docx_resume(
    "neha_gupta_data_science_fresher.docx",
    "Neha Gupta", "neha.gupta@email.com", "+91-9654321098",
    "Final year B.Tech CSE student passionate about data science and machine learning, seeking entry-level opportunities.",
    ["Python", "Machine Learning", "Pandas", "NumPy", "SQL", "Data Analysis", "scikit-learn"],
    [("Data Science Intern", "Startup Analytics Pvt Ltd", "Summer 2025", "Built a churn prediction model using logistic regression and scikit-learn, cleaned and analyzed data using pandas."),
     ("Academic Project", "College Project", "2025", "Built a resume screening system using NLP as final year project.")],
    ["B.Tech in Computer Science (Final Year), Delhi Technological University, 2026 (expected)"],
)

# ---------------- JOB DESCRIPTIONS ----------------

make_jd(
    "jd_data_scientist.txt",
    "Data Scientist", "Amazon India",
    "We are looking for a Data Scientist to join our analytics team. You will build machine learning models "
    "to solve business problems, work with large datasets, and communicate insights to stakeholders.",
    ["Strong experience with Python and libraries like pandas, NumPy, scikit-learn",
     "Experience building and deploying machine learning models",
     "Good understanding of statistics and SQL",
     "Experience with TensorFlow or PyTorch is a plus",
     "Bachelor's degree in Computer Science or related field"],
)

make_jd(
    "jd_backend_developer.txt",
    "Backend Developer (Java)", "Paytm",
    "We are hiring a Backend Developer to design and build scalable microservices for our payments platform. "
    "You will work closely with the platform and infra teams.",
    ["3+ years of experience with Java and Spring Boot",
     "Experience designing REST APIs and microservices",
     "Good understanding of SQL and relational databases",
     "Experience with Docker and Kubernetes preferred",
     "Familiarity with AWS or other cloud platforms"],
)

make_jd(
    "jd_frontend_developer.txt",
    "Frontend Developer (React)", "Swiggy",
    "We are looking for a Frontend Developer to build fast, responsive customer facing applications using React.",
    ["2+ years of experience with JavaScript/TypeScript and React",
     "Strong understanding of HTML, CSS and responsive design",
     "Experience with state management libraries like Redux",
     "Familiarity with Git and modern frontend tooling"],
)

print("Sample data generated successfully.")
print("Resumes:", os.listdir(RESUME_DIR))
print("JDs:", os.listdir(JD_DIR))